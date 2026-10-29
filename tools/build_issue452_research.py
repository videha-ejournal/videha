from docx import Document
from pathlib import Path
import re, json, html, unicodedata
root=Path('.'); doc=Document(r'C:\Users\DELL\Documents\Github_VIDEHA_SADEHA_PDF_DOCX\VIDEHA_SADEHA_DOC_DOCX\VIDEHA_DOCX\Videha 444_452\Videha_452_2.docx')
pat=re.compile(r'^([०-९]+\.[०-९]+)\.?\s*(.*?)\s*[—-]\s*(.+)$')
items=[]
for i,p in enumerate(doc.paragraphs):
 t=' '.join(p.text.split())
 m=pat.match(t)
 if m: items.append((i,m.group(1),m.group(2).strip(),m.group(3).strip()))
out=[]
for n,(start,num,author,title) in enumerate(items):
 end=items[n+1][0] if n+1<len(items) else len(doc.paragraphs)
 body='\n'.join(x.text.strip() for x in doc.paragraphs[start+1:end] if x.text.strip())
 if len(body)<400: continue
 slug=unicodedata.normalize('NFKD',title).encode('ascii','ignore').decode().lower(); slug=re.sub(r'[^a-z0-9]+','-',slug).strip('-') or f'issue452-{n+1}'
 path=f'2026/452/issue-452-{n+1:02d}-{slug[:70]}.htm'; page=root/'research'/path; page.parent.mkdir(parents=True,exist_ok=True)
 page.write_text(f'<!doctype html><html lang="mai"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>{html.escape(title)} · Videha अंक ४५२</title><meta name="citation_title" content="{html.escape(title)}"><meta name="citation_author" content="{html.escape(author)}"><meta name="citation_publication_date" content="2026/10/15"><meta name="citation_journal_title" content="Videha"><meta name="citation_issn" content="2229-547X"><meta name="citation_issue" content="452"><link rel="canonical" href="https://www.videha.co.in/research/{path}"><link rel="stylesheet" href="../../../assets/css/videha-research.css"><script type="application/ld+json">{{"@context":"https://schema.org","@type":"ScholarlyArticle","headline":{json.dumps(title,ensure_ascii=False)},"author":{{"@type":"Person","name":{json.dumps(author,ensure_ascii=False)}}},"datePublished":"2026-10-15","isPartOf":{{"@type":"Periodical","name":"Videha","issn":"2229-547X"}}}}</script></head><body><header class="videha-site-header"><a href="https://videha-ejournal.github.io/videha/">VIDEHA · प्रथम मैथिली पाक्षिक eJournal</a></header><main data-pagefind-body><nav><a href="../../index.htm">शोध-सूची</a> · <a href="https://videha-ejournal.github.io/videha/videha-archive-explorer.html?issue=452">अंक ४५२</a></nav><h1>{html.escape(title)}</h1><p class="article-byline">{html.escape(author)} · Videha अंक ४५२ · 15 October 2026</p><article>{''.join('<p>'+html.escape(x)+'</p>' for x in body.splitlines())}</article></main><footer><a href="https://videha-ejournal.github.io/videha/">VIDEHA Home</a> · <a href="../../index.htm">Research Index</a></footer></body></html>',encoding='utf-8')
 out.append({'title':title,'english_title':None,'authors':[author],'publication_date':'2026-10-15','year':'2026','issue':'452','language':'mai','keywords':['Issue 452','Maithili research'],'classification':'Maithili literary and cultural research','url':'https://videha-ejournal.github.io/videha/research/'+path,'pdf_url':'https://archive.org/download/videha-2026/Videha%20452.pdf','source_url':'https://archive.org/download/videha-2026/Videha%20452.pdf','path':path,'page_start':None,'page_end':None})
inv=root/'research/data/articles.json'; d=json.loads(inv.read_text(encoding='utf-8')); existing={x['path'] for x in d['articles']}; d['articles'] += [x for x in out if x['path'] not in existing]; d['extraction_summary']['published_after_all_overrides']=len(d['articles']); inv.write_text(json.dumps(d,ensure_ascii=False,indent=2),encoding='utf-8'); print('built',len(out),'total',len(d['articles']))
